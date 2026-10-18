"""Backfill CEA's Daily Coal Stock Report per plant, 2017-08-03 .. 2024-06-30, into coal_plants_daily.

The NPP archive keeps one file per report date:
  https://npp.gov.in/public-reports/cea/daily/fuel/DD-MM-YYYY/dailyCoal1-YYYY-MM-DD.{pdf|xls|xlsx}
  2017-08 .. 2018-01-08  PDF only (region/state layout, stock in kt and in days, no receipt/consumption)
  2018-01-09 .. 09-24    nothing (404)
  2018-09-25 .. 2024-06  xlsx (June 2024 is mixed; we try pdf, xls, xlsx)
The sheet layout changed many times (Region/State rows until Nov 2021; A/B/C/D sections from Nov 2021 whose
letters mean different things over time; Utility column from May 2022; bilingual Hindi/English from 2023),
so the parser is header driven: it finds the row holding 'Name of Thermal Power Station', maps each field by
its header text, and reads plant rows below it.

Normalisation (so the rows line up with the 2024-07+ PDFs parsed by scrapers/cea_coal.py):
  section    canonical label by meaning, not by printed letter:
               'A. PLANT HAVING COAL LINKAGES, NO LINKAGES AND COAL BLOCK' (incl. Nov 2021's separate
                  linkage / no-linkage / coal-block sections), 'B. PLANTS DESIGNED ON IMPORTED COAL',
               'C. NOT IN OPERATION', 'D. PLANTS BASE ON WASHERY REJECTS';
             reports without sections (2017 .. Oct 2021) get 'ALL PLANTS (report has no sections)'.
             So `section NOT LIKE 'C.%'` keeps exactly the operating plants in every period.
  state      the state header row above the plant (English part of 'हरियाणा/ Haryana'); from Dec 2021 central
             and private plants sit under group rows 'NTPC', 'NTPC JV', 'DVC', 'IPP' which are stored as the
             state, as cea_coal.py does; reset at each section.
  utility    Utility column when present, else from the 'HPGCL- Total' / 'Total-NTPC' row closing the block.
  plf_pct, stock_pct   0-100 scale (fractions such as 0.16 are multiplied by 100).
  stock_days         stock in days, where the report prints it (2017 .. Nov 2021).
  norm_stock_kt      printed from Oct 2021; earlier = norm_stock_days * daily_req_kt (CEA's own definition).
  stock_pct          printed from Oct 2021; earlier = stock_total_kt / norm_stock_kt * 100.
  receipt_kt / consumption_kt only where printed (Sep 2021 on; consumption from Nov 2021).

    python -m scrapers.cea_coal_hist                      # all NPP dates < 2024-07-01 not yet stored
    python -m scrapers.cea_coal_hist --from 2021-01-01 --to 2021-12-31
    python -m scrapers.cea_coal_hist --file raw/downloads/dailycoal_hist/dailyCoal1-2019-03-15.xlsx --dry
    python -m scrapers.cea_coal_hist --validate
"""
import argparse
import re
import sqlite3
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import requests
import urllib3

from store import connect, log

urllib3.disable_warnings()
ROOT = Path(__file__).resolve().parent.parent
DL = ROOT / "raw" / "downloads" / "dailycoal_hist"
URL = "https://npp.gov.in/public-reports/cea/daily/fuel/{d:%d-%m-%Y}/dailyCoal1-{d:%Y-%m-%d}.{ext}"
INDEX = "https://npp.gov.in/dailyCoalReports"
UA = {"User-Agent": "Mozilla/5.0 (reclaimchennai power map; +https://power.reclaimchennai.city)"}
FIRST, LAST = date(2017, 8, 3), date(2024, 6, 30)

SEC_A = "A. PLANT HAVING COAL LINKAGES, NO LINKAGES AND COAL BLOCK"
SEC_B = "B. PLANTS DESIGNED ON IMPORTED COAL"
SEC_C = "C. NOT IN OPERATION"
SEC_D = "D. PLANTS BASE ON WASHERY REJECTS"
SEC_ALL = "ALL PLANTS (report has no sections)"
REGIONS = {"northern", "western", "southern", "eastern", "north eastern", "north-eastern", "northeastern"}
DATE_RE = re.compile(r"as\s+on\s*[:\-]?\s*\(?\s*(\d{1,2})\s*[.\-/]\s*(\d{1,2})\s*[.\-/]\s*(\d{4})", re.I)
NUM_RE = re.compile(r"^-?[\d,]*\.?\d+$")
FIELDS = ("plant", "utility", "state", "section", "mode", "capacity_mw", "plf_pct", "norm_stock_days", "daily_req_kt",
          "norm_stock_kt", "stock_indigenous_kt", "stock_import_kt", "stock_total_kt", "stock_pct", "critical",
          "receipt_kt", "consumption_kt", "remarks", "stock_days")


def _t(v):
    """cell -> single-spaced text ('' for None)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _f(v):
    if v is None or isinstance(v, (datetime, date)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    v = str(v).replace(",", "").strip().rstrip("%").strip()
    try:
        return float(v)
    except ValueError:
        return None


def english(s):
    """'हरियाणा/ Haryana' -> 'Haryana'; plain English is returned unchanged."""
    if re.search(r"[ऀ-ॿ]", s) and "/" in s:
        s = s.rsplit("/", 1)[1]
    return s.strip()


def canon_section(s):
    u = s.upper()
    if "NOT IN OPERATION" in u or "वर्तमान" in s:
        return SEC_C
    if "IMPORT" in u or "आयातित" in s:
        return SEC_B
    if "WASHERY" in u or "वाशरी" in s:
        return SEC_D
    if "LINKAGE" in u or "COAL BLOCK" in u or "लिंकेज" in s:
        return SEC_A
    return None


# ---------------------------------------------------------------------------------------------- header mapping
def header_map(rows, i):
    """Column index per field from header row i (+ up to 3 sub-header rows below it)."""
    head = [_t(c).lower() for c in rows[i]]
    head = [re.sub(r"(\w)- ?(\w)", r"\1\2", h) for h in head]
    squash = [h.replace(" ", "") for h in head]                   # PDF headers wrap mid-word: 'Requireme nt'
    cols = {}

    def put(k, j):
        cols.setdefault(k, j)

    for j, t in enumerate(head):
        if not t:
            continue
        if j == 0 or t in ("state", "region/state", "region/ state") or t.startswith("sl no"):
            put("label", j)
        if t.startswith("region") or t.endswith("state"):
            put("label", j)
        if t.startswith("mode of"):
            put("mode", j)
        elif t.startswith("name of thermal"):
            put("plant", j)
        elif t.startswith("utility"):
            put("utility", j)
        elif t.startswith("capacity"):
            put("capacity_mw", j)
        elif t.startswith("normative"):
            put("norm_stock_days" if "day" in t else "norm_stock_kt", j)
        elif squash[j].startswith(("requirementfortheday", "dailyrequirement")):
            put("daily_req_kt", j)
        elif "plf" in t and "spcc" not in t:
            if "plf_pct" not in cols:
                cols["plf_pct"] = j
                cols["_plf_frac"] = "tentative" in t
        elif ("%" in t or "percentage" in t) and "actual stock" in t:
            put("stock_pct", j)
        elif t.startswith("actual stock") or (t.startswith("indigenous/") and "stock" in t):
            put("_stock", j)
            if "indigenous/" in t:                             # Mar 2022: one merged 'Indigenous/ Import' column
                cols["_single"] = True
        elif t.startswith("critical"):
            put("critical", j)
        elif t.startswith("receipt"):
            put("receipt_kt", j)
        elif t.startswith("consumption"):
            put("consumption_kt", j)
        elif t.startswith("reasons"):
            put("remarks", j)
    if "plant" not in cols or "_stock" not in cols:
        return None, i
    # sub-headers: Indigenous / Import / Total, 'In Days'; skip them and the '1 2 3 ...' column-number row
    k = i + 1
    while k < len(rows) and k <= i + 4:
        sub = [_t(c).lower() for c in rows[k]]
        filled = [c for c in sub if c]
        if not filled:
            k += 1
            continue
        nums = [c for c in filled if NUM_RE.match(c)]
        if len(nums) >= 6 and len(nums) >= len(filled) - 1:          # the column-number row
            k += 1
            continue
        if any(c.startswith(("indigenou", "import", "total", "in '000", "in day", "in ('000")) for c in filled):
            for j, c in enumerate(sub):
                if c.startswith("indigenou"):
                    put("stock_indigenous_kt", j)
                elif c.startswith("import"):
                    put("stock_import_kt", j)
                elif c == "total":
                    put("stock_total_kt", j)
                elif c.startswith("in day"):
                    put("stock_days", j)
            k += 1
            continue
        break
    s = cols.pop("_stock")
    if cols.pop("_single", False) and "stock_indigenous_kt" not in cols:
        cols["stock_total_kt"] = s                             # only the total is printed
        cols.setdefault("label", 0)
        return cols, k
    cols.setdefault("stock_indigenous_kt", s)
    cols.setdefault("stock_import_kt", cols["stock_indigenous_kt"] + 1)
    cols.setdefault("stock_total_kt", cols["stock_indigenous_kt"] + 2)
    cols.setdefault("label", 0)
    return cols, k


def find_date(rows):
    for r in rows[:12]:
        for c in r:
            if isinstance(c, (datetime, date)):
                continue
            m = DATE_RE.search(_t(c))
            if m:
                d, mo, y = map(int, m.groups())
                return date(y, mo, d).isoformat()
    return None


# ---------------------------------------------------------------------------------------------- row parser
def parse_rows(rows, sectioned_default=None):
    """rows: list of cell lists (one sheet, or all PDF tables concatenated). Returns (date, records, notes)."""
    rdate = find_date(rows)
    cols, section, state, pending, out, notes = None, sectioned_default, None, [], [], []
    has_sections = False
    i = 0
    while i < len(rows):
        raw = list(rows[i])
        cells = [_t(c) for c in raw]
        if not any(cells):
            i += 1
            continue
        if any(c.lower().startswith("name of thermal") for c in cells):
            hm, nxt = header_map(rows, i)
            if hm:
                cols = hm
                i = nxt
                continue
        if cols is None:
            i += 1
            continue
        low = " ".join(cells).lower()
        if "report summary" in low or "grand total" in low or "all india total" in low:
            break
        li, pi = cols["label"], cols["plant"]
        label = cells[li] if li < len(cells) else ""
        name = cells[pi] if pi < len(cells) else ""
        filled = [c for c in cells if c]
        # section heading 'A. PLANTS HAVING ...'
        if re.match(r"^[A-F]\s*\.\s*\S", label) and len(filled) <= 2 and not NUM_RE.match(label):
            sec = canon_section(label)
            if sec is None:
                notes.append(f"unknown section {label[:60]!r}")
                sec = label
            section, state, pending, has_sections = sec, None, [], True
            i += 1
            continue
        # total / subtotal rows close a utility block
        tot = next((c for c in cells[:pi + 1] if "total" in c.lower()), None)
        if tot:
            m = re.match(r"^(.*?)\s*-\s*total$", tot, re.I) or re.match(r"^total\s*-\s*(.*)$", tot, re.I)
            u = m.group(1).strip() if m else None
            if u and not re.match(r"^[A-F](\b|/)", u) and "&" not in u.replace("APGENCO & APPDCL", ""):
                for r in pending:
                    r["utility"] = r["utility"] or u
            pending = []
            i += 1
            continue
        # region / state / group label on a row of its own
        if label and not NUM_RE.match(label) and not name and len(filled) <= 2:
            t = english(label)
            if t.lower() in REGIONS or t.lower() in ("state", "region/state"):
                pass
            else:
                state = t
            i += 1
            continue
        if not name or NUM_RE.match(name) or name.lower() in ("state",):
            i += 1
            continue
        rec = {k: None for k in FIELDS}
        rec.update(plant=name, state=state, section=section or SEC_ALL)
        for k, j in cols.items():
            if k.startswith("_") or k in ("label", "plant") or j >= len(raw):
                continue
            if k in ("mode", "utility", "remarks"):
                rec[k] = cells[j] or None
            elif k == "critical":
                rec[k] = 1 if "*" in cells[j] else 0
            else:
                rec[k] = _f(raw[j])
        rec["critical"] = rec["critical"] or 0
        if rec["capacity_mw"] is None and rec["stock_total_kt"] is None:
            i += 1
            continue
        out.append(rec)
        pending.append(rec)
        i += 1
    normalise(out, cols or {})
    if not has_sections:
        for r in out:
            r["section"] = SEC_ALL
    return rdate, out, notes


def normalise(recs, cols):
    if not recs:
        return
    plf = [r["plf_pct"] for r in recs if r["plf_pct"] is not None]
    if plf and (cols.get("_plf_frac") or max(plf) <= 1.5):
        for r in recs:
            if r["plf_pct"] is not None:
                r["plf_pct"] = round(r["plf_pct"] * 100, 2)
    for r in recs:
        if r["stock_total_kt"] is None and r["stock_indigenous_kt"] is not None:
            r["stock_total_kt"] = (r["stock_indigenous_kt"] or 0) + (r["stock_import_kt"] or 0)
        if r["norm_stock_kt"] is None and r["norm_stock_days"] is not None and r["daily_req_kt"] is not None:
            r["norm_stock_kt"] = round(r["norm_stock_days"] * r["daily_req_kt"], 3)
    # printed % of normative: fraction (0.16) or percent (16)? decide against stock/normative
    pairs = [(r["stock_pct"], r["stock_total_kt"] / r["norm_stock_kt"]) for r in recs
             if r["stock_pct"] is not None and r["stock_total_kt"] and r["norm_stock_kt"]]
    if pairs:
        frac = sum(1 for p, q in pairs if abs(p - q) < abs(p - 100 * q))
        if frac > len(pairs) / 2:
            for r in recs:
                if r["stock_pct"] is not None:
                    r["stock_pct"] = round(r["stock_pct"] * 100, 1)
    for r in recs:
        if r["stock_pct"] is None and r["stock_total_kt"] is not None and r["norm_stock_kt"]:
            r["stock_pct"] = round(100 * r["stock_total_kt"] / r["norm_stock_kt"], 1)


# ---------------------------------------------------------------------------------------------- file readers
def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = next((n for n in wb.sheetnames if n.strip().lower() == "dailycoalreport"), wb.sheetnames[0])
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    wb.close()
    return rows


def read_xls(path):
    import xlrd
    wb = xlrd.open_workbook(str(path))
    names = wb.sheet_names()
    ws = wb.sheet_by_name(next((n for n in names if n.strip().lower() == "dailycoalreport"), names[0]))
    return [ws.row_values(i) for i in range(ws.nrows)]


def read_pdf(path):
    import pdfplumber
    rows = []
    with pdfplumber.open(path) as pdf:
        text = pdf.pages[0].extract_text() or ""
        m = DATE_RE.search(text)
        if m:
            rows.append([f"(As on {m.group(1)}.{m.group(2)}.{m.group(3)})"])
        for page in pdf.pages:
            for table in page.extract_tables():
                rows.extend([list(r) for r in table])
    return rows


def parse_file(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        rows = read_pdf(path)
        rd = find_date(rows)
        if rd is None or rd >= "2024-01-01":
            from scrapers import cea_coal                            # 2024 PDF layout: reuse the live parser
            d, recs = cea_coal.parse_pdf(path)
            for r in recs:
                r.setdefault("stock_days", None)
                for k in FIELDS:
                    r.setdefault(k, None)
            return d, recs, ["parsed with cea_coal.parse_pdf"]
        return parse_rows(rows)
    if ext == ".xls":
        return parse_rows(read_xls(path))
    return parse_rows(read_xlsx(path))


# ---------------------------------------------------------------------------------------------- storage
def ensure_columns(con):
    have = {r[1] for r in con.execute("PRAGMA table_info(coal_plants_daily)")}
    if "stock_days" not in have:
        con.execute("ALTER TABLE coal_plants_daily ADD COLUMN stock_days REAL")
        con.commit()


def _rows(d, recs):
    seen, rows = set(), []
    for r in recs:
        name = r["plant"]
        if name in seen:                                      # same name twice in one report: keep both, suffix state
            name = f"{name} ({r['state'] or r['section'][:1]})"
            if name in seen:
                continue
        seen.add(name)
        rows.append(dict(r, plant=name, date=d))
    return rows


INSERT = (f"INSERT OR IGNORE INTO coal_plants_daily (date, {', '.join(FIELDS)}) "
          f"VALUES (:date, {', '.join(':' + k for k in FIELDS)})")


def _retry(fn, tries=10):
    for k in range(tries):                                    # other collectors/backfills write the same db
        try:
            return fn()
        except sqlite3.OperationalError as e:
            if "locked" not in str(e) or k == tries - 1:
                raise
            time.sleep(5 + 5 * k)


def save(d, recs, replace=False):
    """Insert one report day. replace=False skips a date already stored; True swaps its rows in one transaction."""
    if not (FIRST.isoformat() <= d <= LAST.isoformat()):
        raise ValueError(f"refusing to write {d}: outside {FIRST}..{LAST}")
    rows = _rows(d, recs)

    def go():
        con = connect()
        try:
            ensure_columns(con)
            with con:
                if replace:
                    con.execute("DELETE FROM coal_plants_daily WHERE date=?", (d,))
                elif con.execute("SELECT 1 FROM coal_plants_daily WHERE date=? LIMIT 1", (d,)).fetchone():
                    return 0                                  # stored meanwhile (e.g. by backfill_coal for June 2024)
                con.executemany(INSERT, rows)
            return len(rows)
        finally:
            con.close()
    return _retry(go)


def reparse(start, end):
    """Re-parse cached files for stored dates in [start, end] and swap in the new rows (after a parser fix)."""
    con = connect()
    have = sorted(r[0] for r in con.execute("SELECT DISTINCT date FROM coal_plants_daily WHERE date BETWEEN ? AND ?",
                                           (max(start, FIRST).isoformat(), min(end, LAST).isoformat())))
    con.close()
    n = 0
    for d in have:
        p = cached(d)
        if p is None or (p.suffix == ".pdf" and d >= "2024-01-01"):
            continue                                          # 2024 PDFs belong to backfill_coal / cea_coal
        rd, recs, notes = parse_file(p)
        if not recs:
            print(d, "no rows on reparse; kept old rows", flush=True)
            continue
        print(d, len(recs), save(d, recs, replace=True), f"{national(recs) / 1000:.2f}Mt", flush=True)
        n += 1
    print(f"re-parsed {n} dates")


def stored_dates():
    con = connect()
    try:
        return {r[0] for r in con.execute("SELECT DISTINCT date FROM coal_plants_daily")}
    finally:
        con.close()


# ---------------------------------------------------------------------------------------------- fetching
def npp_dates(session):
    """All report dates NPP lists for the daily coal report (the page embeds `var date = "YYYY-MM-DD"` per entry)."""
    cache = DL / "npp_coal_dates.txt"
    try:
        r = session.get(INDEX, timeout=120, verify=False)
        ds = sorted(set(re.findall(r'var date = "(\d{4}-\d{2}-\d{2})"', r.text)))
        if len(ds) > 500:
            cache.write_text("\n".join(ds) + "\n")
            return ds
        print(f"index gave {len(ds)} dates", file=sys.stderr)
    except Exception as e:
        print("index fetch failed:", e, file=sys.stderr)
    if cache.exists():
        return cache.read_text().split()
    d, out = FIRST, []                                         # fall back to walking every day
    while d <= LAST:
        out.append(d.isoformat())
        d += timedelta(days=1)
    return out


def cached(d):
    for ext in ("pdf", "xls", "xlsx"):
        p = DL / f"dailyCoal1-{d}.{ext}"
        if p.exists() and p.stat().st_size > 2000:
            return p
    return None


MAGIC = {"pdf": b"%PDF", "xls": b"\xd0\xcf\x11\xe0", "xlsx": b"PK"}


def fetch(session, d):
    p = cached(d)
    if p:
        return p, 0
    dd = date.fromisoformat(d)
    if dd <= date(2018, 1, 8):
        order = ("pdf",)
    elif dd >= date(2024, 5, 1):
        order = ("pdf", "xls", "xlsx")
    else:
        order = ("xlsx", "xls", "pdf")
    n = 0
    for ext in order:
        r = session.get(URL.format(d=dd, ext=ext), timeout=120, verify=False)
        n += 1
        time.sleep(1.0)
        if r.status_code == 200 and r.content.startswith(MAGIC[ext]):
            p = DL / f"dailyCoal1-{d}.{ext}"
            p.write_bytes(r.content)
            return p, n
    return None, n


def national(recs):
    return sum(r["stock_total_kt"] or 0 for r in recs if not (r["section"] or "").startswith("C."))


def run(start, end, dry=False):
    DL.mkdir(parents=True, exist_ok=True)
    s = requests.Session()
    s.headers.update(UA)
    dates = [d for d in npp_dates(s) if max(start, FIRST).isoformat() <= d <= min(end, LAST).isoformat()]
    have = stored_dates()
    todo = [d for d in dates if d not in have]
    print(f"{len(dates)} NPP dates in range, {len(todo)} not stored", flush=True)
    ok = miss = bad = 0
    for d in todo:
        try:
            p, _ = fetch(s, d)
            if p is None:
                miss += 1
                print(d, "missing", flush=True)
                continue
            rd, recs, notes = parse_file(p)
            if not recs:
                bad += 1
                print(d, "error no rows", p.name, flush=True)
                continue
            use = d
            if rd != d:
                print(d, f"WARN sheet date {rd} != file date", flush=True)
                if rd and rd not in have and rd not in dates and FIRST.isoformat() <= rd <= LAST.isoformat():
                    use = rd                                   # misfiled report that has no file of its own
            n = 0 if dry else save(use, recs)
            have.add(use)
            ok += 1
            print(d, use, p.suffix, len(recs), n, f"{national(recs) / 1000:.2f}Mt", *notes, flush=True)
        except Exception as e:
            bad += 1
            print(d, "error", repr(e)[:300], flush=True)
    msg = f"{ok} days stored, {miss} missing, {bad} errors ({start}..{end})"
    print(msg)
    if not dry:
        log("cea_coal_hist", "ok" if not bad else "partial", msg)


EXPECTED = {"2019-03-16": 28.96, "2019-11-01": 19.27, "2020-06-01": 49.50, "2021-10-01": 7.96,
            "2022-04-29": 22.04, "2023-10-02": 23.51, "2024-06-01": 47.63}


def validate():
    con = connect()
    print("date        plants  ours_Mt  expected")
    for d, exp in EXPECTED.items():
        n, s = con.execute("SELECT count(*), sum(stock_total_kt) FROM coal_plants_daily "
                           "WHERE date=? AND section NOT LIKE 'C.%'", (d,)).fetchone()
        print(f"{d}  {n:6d}  {(s or 0) / 1000:7.2f}  {exp:7.2f}")
    con.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--from", dest="start", default=FIRST.isoformat())
    ap.add_argument("--to", dest="end", default=LAST.isoformat())
    ap.add_argument("--file")
    ap.add_argument("--dry", action="store_true")
    ap.add_argument("--validate", action="store_true")
    ap.add_argument("--reparse", action="store_true", help="re-parse cached files for stored dates in range")
    a = ap.parse_args()
    if a.validate:
        return validate()
    if a.reparse:
        return reparse(date.fromisoformat(a.start), date.fromisoformat(a.end))
    if a.file:
        rd, recs, notes = parse_file(a.file)
        print(rd, len(recs), f"{national(recs) / 1000:.2f}Mt", notes)
        for r in recs[:3] + recs[-3:]:
            print(r)
        if not a.dry:
            print(save(rd, recs))
        return
    run(date.fromisoformat(a.start), date.fromisoformat(a.end), a.dry)


if __name__ == "__main__":
    main()
