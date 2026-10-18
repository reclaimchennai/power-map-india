"""Gas feedstock of India's gas power stations, monthly.

1) PPAC sectoral natural-gas consumption (MMSCM in the month), https://ppac.gov.in/natural-gas/sectoral-consumption
   <ts>_2_NG-H_Sectoral_Consumption.xlsx   one sheet per FY, FY2015-16 onwards (monthly, by sector, Domestic/RLNG/Total)
   <ts>_4_NG-C-Sectoral-Consumption.xlsx   current FY to date, plus one state x sector sheet per month
   (file names carry an upload timestamp, so the page is read for the current links). Series in coal_supply
   (scrapers.coal_supply.put, a generic series table despite its name), period 'YYYY-MM', unit MMSCM:
     gas_power_m:Domestic | :RLNG | :Total          gas consumed by the power sector
     gas_sector_m:<Sector>                          total per sector (Power, CGD, Fertilizer, Refinery, ...)
     gas_total_m:Domestic | :RLNG | :Total          all sectors
     gas_power_state_m:<STATE>                      power-sector gas by state (current-FY file only)
   From FY2020-21 PPAC adds ONGC/OIL/DGH internal consumption (incl. captive power), a small step in the series.

2) CEA Fuel Management Division "Monthly report on fuel supply/consumption for gas based power stations",
   listed by the CEA WordPress AJAX archive (action=monthly_archive_report, reportType=coal: the gas report sits
   in the coal/fuel list); 26 more months exist on the server unlisted and are fetched from UNLISTED (found by
   probing file names). Coverage: 2012-01, 03, 04, 05, 2015-06, 2016-10 -> 2021-01, 2021-04 -> 2022-06,
   2022-08 -> 2025-05. Missing: 2021-02/03, 2022-07, anything after 2025-05; 2016-09 is a scanned image (skipped).
   Table gas_plants_monthly, one row per plant and month (MMSCMD = million standard m3 per day, month average).
   Layouts, told apart by the number of cells in the GRAND TOTAL row (LAYOUTS):
      7  2012             requirement at 90% PLF, allotted, consumed (total incl. RLNG), naphtha, HSD, gen loss
     15  2015-06          allotted/consumed split APM / non-APM / KG-D6 / RLNG (LT, spot), gen loss
     17  2016-10..2017-10 same with RLNG LT / e-bid / spot, gen loss
     16  2017-11..2018-08 as 17 without gen loss
     13  2018-09..2023-12 domestic (APM+non-APM+PMT together, KG-D6/auctioned) and RLNG (LT, spot); no gen loss
     14  2024-01..        as 13 plus a separate RLNG total; from 2024-04 plants grouped grid / isolated (supply P/I)
   9, 10 2021-10, 2024-09 consumption only (no allotment columns)
   Cells left blank in the PDF are placed by checking the row's own totals (realign); totals are rebuilt from
   parts when blank. Plant names are as printed and drift between layouts ('NTPC, ANTA CCPP' -> 'ANTA CCPP').
   National series in coal_supply (period 'YYYY-MM'): gas_cea_m:Consumed | :Allotted | :Required (MMSCMD),
   gas_cea_m:Consumed_RLNG | :Consumed_Domestic (MMSCMD), gas_gen_m:Total (MU, generation of the listed stations),
   gas_cea_m:Allotted_Domestic | :Allotted_RLNG, gas_genloss_m:Total (MU lost to gas shortage; 2012, 2015-06,
   2016-10..2017-10 only), gas_cea_cap_m:Total (MW), all from each report's GRAND TOTAL row.
   Annual reports (xlsx, FY2022-23 and FY2024-25: yearly-average MMSCMD, annual MU) go to gas_plants_annual.

NPP (npp.gov.in/publishedReports) lists a "Monthly Gas Supply" report but it is greyed out (no file), and no
monthGas*-YYYY-Mon.xls/pdf exists under public-reports/cea/monthly/fuel/, so the CEA series ends May 2025.

    python -m scrapers.gas_power              # refresh: PPAC files + CEA archive for the last 4 months
    python -m scrapers.gas_power --backfill   # everything (cached downloads reused)
Downloads are cached under raw/gas/.
"""
import re
import sys
import time
from itertools import combinations
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pdfplumber
import requests
import urllib3

from scrapers.coal_supply import put
from store import connect, log

urllib3.disable_warnings()
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"}
RAW = Path(__file__).resolve().parent.parent / "raw" / "gas"
PPAC_PAGE = "https://ppac.gov.in/natural-gas/sectoral-consumption"
PPAC_FILE = "https://ppac.gov.in/uploads/page-images/{}"
CEA_AJAX = "https://cea.nic.in/wp-admin/admin-ajax.php"
CEA_REFERER = "https://cea.nic.in/fuel-reports/?lang=en"
CEA_ANNUAL_REFERER = "https://cea.nic.in/annual-fuel-consumption-report/?lang=en"
FIRST_ARCHIVE = (2012, 1)
SRC_PPAC = "PPAC sectoral natural gas consumption"
SRC_CEA = "CEA monthly fuel supply/consumption report for gas based power stations"

SCHEMA = """
CREATE TABLE IF NOT EXISTS gas_plants_monthly (
    month TEXT NOT NULL,              -- 'YYYY-MM'
    plant TEXT NOT NULL,              -- as named in the report (names drift between layouts)
    sector TEXT,                      -- Central | State | Private
    supply TEXT,                      -- P (pipeline/grid) | I (isolated field)
    state TEXT,
    capacity_mw REAL,
    gen_mu REAL,                      -- generation in the month (tentative)
    gas_required_mmscmd REAL,         -- normative requirement at 90% PLF (2012/2015 layout only)
    allotted_apm REAL,                -- APM (from 2018-09: APM + non-APM + PMT)
    allotted_nonapm REAL,             -- non-APM (2015-06..2018-08 only)
    allotted_kgd6 REAL,
    allotted_domestic REAL,           -- APM + non-APM + KG-D6
    allotted_rlng REAL,               -- RLNG long term (+ e-bid 2016-10..2018-08)
    gas_allotted_mmscmd REAL,         -- total allotted
    consumed_apm REAL,                -- as allotted_apm
    consumed_nonapm REAL,
    consumed_kgd6 REAL,
    consumed_domestic REAL,
    consumed_rlng REAL,               -- all RLNG: LT + e-bid + spot
    consumed_rlng_spot REAL,
    gas_consumed_mmscmd REAL,         -- total consumed/supplied
    naphtha_kl REAL,
    hsd_kl REAL,
    gen_loss_mu REAL,                 -- generation lost to short gas supply (2012, 2015-06, 2016-10..2017-10)
    remarks TEXT,
    source_file TEXT,
    PRIMARY KEY (month, plant)
);
CREATE TABLE IF NOT EXISTS gas_plants_annual (
    fy TEXT NOT NULL,                 -- '2024-25'
    plant TEXT NOT NULL,
    sector TEXT, supply TEXT, state TEXT, capacity_mw REAL, gen_mu REAL,
    allotted_apm REAL, allotted_kgd6 REAL, allotted_domestic REAL, allotted_rlng REAL, gas_allotted_mmscmd REAL,
    consumed_apm REAL, consumed_kgd6 REAL, consumed_domestic REAL, consumed_rlng_lt REAL, consumed_rlng_spot REAL,
    consumed_rlng REAL, gas_consumed_mmscmd REAL, naphtha_kl REAL, hsd_kl REAL, source_file TEXT,
    PRIMARY KEY (fy, plant)
);
"""
COLS = ["month", "plant", "sector", "supply", "state", "capacity_mw", "gen_mu", "gas_required_mmscmd",
        "allotted_apm", "allotted_nonapm", "allotted_kgd6", "allotted_domestic", "allotted_rlng", "gas_allotted_mmscmd",
        "consumed_apm", "consumed_nonapm", "consumed_kgd6", "consumed_domestic", "consumed_rlng", "consumed_rlng_spot",
        "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl", "gen_loss_mu", "remarks", "source_file"]
ACOLS = ["fy", "plant", "sector", "supply", "state", "capacity_mw", "gen_mu", "allotted_apm", "allotted_kgd6",
         "allotted_domestic", "allotted_rlng", "gas_allotted_mmscmd", "consumed_apm", "consumed_kgd6", "consumed_domestic",
         "consumed_rlng_lt", "consumed_rlng_spot", "consumed_rlng", "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl", "source_file"]

_last = [0.0]


def fetch(url, dest=None, post=None, referer=None, refresh=False):
    """GET (or POST) politely (>= 1 s apart); cache to dest when given. Returns bytes or None on 404."""
    if dest is not None and dest.exists() and dest.stat().st_size > 0 and not refresh:
        return dest.read_bytes()
    wait = 1.0 - (time.time() - _last[0])
    if wait > 0:
        time.sleep(wait)
    headers = dict(UA)
    if referer:
        headers["Referer"] = referer
    for attempt in range(3):
        try:
            r = (requests.post(url, data=post, headers=headers, timeout=120, verify=False) if post
                 else requests.get(url, headers=headers, timeout=120, verify=False))
            break
        except requests.RequestException:
            if attempt == 2:
                raise
            time.sleep(5)
    _last[0] = time.time()
    if r.status_code == 404:
        return None
    r.raise_for_status()
    if dest is not None:
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(r.content)
    return r.content


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").strip()
    if s in ("", "-", "--", "NA", "DNR"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


# ----------------------------------------------------------------------------------------------- PPAC

SECTOR = {"power": "Power", "cgd": "CGD", "refinery": "Refinery", "i?c for p/l system": "Pipeline I/C",
          "i/c for p/l system": "Pipeline I/C", "agriculture(tea plantation)": "Agriculture", "industrial": "Industrial",
          "manufacturing": "Manufacturing", "other/misc": "Other/Misc", "fertilizer": "Fertilizer",
          "petrochemical": "Petrochemical", "lpg shrinkage": "LPG Shrinkage", "sponge iron/steel": "Sponge Iron/Steel",
          "total": "Total"}
MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]


def month_of(v):
    """'YYYY-MM' from a header cell (datetime, 'YYYY-MM', 'Apr-26', 'June-26'); None otherwise."""
    if isinstance(v, datetime):
        return f"{v:%Y-%m}"
    s = str(v or "").strip()
    m = re.fullmatch(r"(\d{4})-(\d{2})(?:-\d{2})?(?: 00:00:00)?", s)
    if m:
        return f"{m[1]}-{m[2]}"
    m = re.match(r"([A-Za-z]{3})[a-z]*[-' ,]*(\d{2,4})\b", s)
    if m and m[1].lower() in MONTHS:
        y = int(m[2]) + (2000 if len(m[2]) == 2 else 0)
        return f"{y}-{MONTHS.index(m[1].lower()) + 1:02d}"
    return None


def ppac_sheet(ws):
    """Monthly sector rows of one FY sheet -> {(sector, 'YYYY-MM'): (dom, rlng, total)}."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = next((i for i, r in enumerate(rows) if sum(1 for v in r if month_of(v)) >= 3), None)
    if hdr is None:
        return {}
    cols = [(j, month_of(v)) for j, v in enumerate(rows[hdr]) if month_of(v)]
    out = {}
    for r in rows[hdr + 2:]:
        label = str(r[0] or "").strip().lower()
        if label not in SECTOR:
            continue
        for j, mon in cols:
            dom, rl, tot = (num(r[j + k]) if j + k < len(r) else None for k in range(3))
            if tot is None and dom is None and rl is None:
                continue
            if tot is None:
                tot = (dom or 0) + (rl or 0)
            out[(SECTOR[label], mon)] = (dom, rl, tot)
    # drop months that are not yet filled in (all-zero or missing total row)
    empty = {mon for (s, mon), v in out.items() if s == "Total" and not v[2]}
    have_total = {mon for (s, mon) in out if s == "Total"}
    return {k: v for k, v in out.items() if k[1] not in empty and k[1] in have_total}


def ppac_statewise(ws, mon):
    """Power-sector total by state from a '<Mon-YY> Statewise' sheet."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for i, r in enumerate(rows):
        pcol = next((j for j, v in enumerate(r) if str(v or "").strip().upper() == "POWER"), None)
        if pcol is None:
            continue
        scol = next((j for j, v in enumerate(rows[i + 1]) if str(v or "").strip().lower() == "state"), None)
        out = []
        for rr in rows[i + 2:]:
            st = str(rr[scol] or "").strip() if scol is not None else ""
            if not st or st.lower().startswith(("grand total", "*")):
                continue
            v = num(rr[pcol + 2]) if pcol + 2 < len(rr) else None
            if v is not None:
                out.append((f"gas_power_state_m:{st.upper()}", mon, round(v, 4), "MMSCM", SRC_PPAC))
        return out
    return []


def ppac(refresh=True):
    page = fetch(PPAC_PAGE, RAW / "ppac" / "sectoral-consumption.html", refresh=refresh).decode("utf8", "ignore")
    names = sorted(set(re.findall(r"(\d{9,11}_[24]_NG-[HC][-_]Sectoral[-_]Consumption\.xlsx)", page)))
    if not names:
        raise RuntimeError("PPAC page: no sectoral consumption links found")
    # historical file first so the current-FY file wins on overlap
    names.sort(key=lambda n: ("_4_" in n, n))
    data, state_rows, files = {}, [], []
    for n in names:
        path = RAW / "ppac" / n
        if fetch(PPAC_FILE.format(n), path) is None:
            log("gas_power", "error", f"PPAC 404 {n}")
            continue
        files.append(n)
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            if "statewise" in ws.title.lower():
                mon = month_of(ws.title)
                if mon:
                    state_rows += ppac_statewise(ws, mon)
            elif ws.title.strip() != "NG-H Sectoral":      # the FY-summary sheet
                data.update(ppac_sheet(ws))
    rows = []
    for (sec, mon), (dom, rl, tot) in sorted(data.items()):
        rows.append((f"gas_sector_m:{sec}" if sec != "Total" else "gas_total_m:Total", mon, round(tot, 4), "MMSCM", SRC_PPAC))
        if sec in ("Power", "Total"):
            base = "gas_power_m" if sec == "Power" else "gas_total_m"
            if sec == "Power":
                rows.append((f"{base}:Total", mon, round(tot, 4), "MMSCM", SRC_PPAC))
            for part, v in (("Domestic", dom), ("RLNG", rl)):
                if v is not None:
                    rows.append((f"{base}:{part}", mon, round(v, 4), "MMSCM", SRC_PPAC))
    rows += state_rows
    put(rows)
    months = sorted({m for (s, m) in data if s == "Power"})
    fy16 = sum(v[2] for (s, m), v in data.items() if s == "Power" and "2015-04" <= m <= "2016-03")
    msg = f"PPAC {len(rows)} rows, power {months[0]}..{months[-1]} ({len(months)} months); FY2015-16 power {fy16:,.0f} MMSCM; files {files}"
    log("gas_power", "ok", msg)
    return msg


# ----------------------------------------------------------------------------------------------- CEA listing

def months(a, b):
    y, m = a
    while (y, m) <= b:
        yield y, m
        y, m = (y, m + 1) if m < 12 else (y + 1, 1)


def cea_links(y, m, refresh=False):
    """Gas-report links (pdf/xlsx) the CEA monthly archive lists for a month."""
    body = fetch(CEA_AJAX, RAW / "cea_ajax" / f"{y}-{m:02d}.html", referer=CEA_REFERER, refresh=refresh,
                 post={"action": "monthly_archive_report", "selMonthYear": f"{y}-{m:02d}", "reportType": "coal"})
    links = re.findall(r'href="(https://cea\.nic\.in/wp-content/uploads/[^"]+\.(?:pdf|xlsx?))"', (body or b"").decode("utf8", "ignore"), re.I)
    out = []
    for u in links:
        if re.search(r"gas", u.rsplit("/", 1)[-1], re.I) and u not in out:
            out.append(u)
    return out


# Reports that exist on cea.nic.in but that the archive listing does not show (found by probing file names).
UNLISTED = {
    "2021-04": "coal/2021/04/Monthly_Gas_Report_April_2021.pdf",
    "2021-06": "coal/2021/08/Monthly_Gas_Report_Jun_21.pdf",
    "2021-08": "coal/2021/08/Monthly_Gas_Report_August_21.pdf",
    "2021-09": "coal/2021/09/Monthly_Gas_Report_September_21.pdf",
    "2021-10": "coal/2021/10/Monthly_Gas_Report_October_21.pdf",
    "2021-11": "coal/2021/11/Monthly_Gas_Report_November_21.pdf",
    "2021-12": "coal/2022/01/Monthly_Gas_Report_December_21.pdf",
    "2022-01": "coal/2022/01/Monthly_Gas_Report_January_22.pdf",
    "2022-02": "coal/2022/02/Monthly_Gas_Report_February_22.pdf",
    "2022-03": "coal/2022/03/Monthly_Gas_Report_March_22.pdf",
    "2022-05": "coal/2022/05/Monthly_Gas_Report_May_22.pdf",
    "2022-06": "coal/2022/07/Monthly_Gas_Report_June_22.pdf",
    "2022-08": "coal/2022/09/Monthly_Gas_Report_Aug_22.pdf",
    "2022-09": "coal/2022/09/Monthly_Gas_Report_September_22.pdf",
    "2022-10": "coal/2022/11/Monthly_Gas_Report_October_22.pdf",
    "2022-11": "coal/2022/11/Monthly_Gas_Report_November_22_upload.pdf",
    "2022-12": "coal/2023/01/Monthly_Gas_Report_December_22_upload.pdf",
    "2023-01": "coal/2023/02/Monthly_Gas_Report_Jan_23.pdf",
    "2023-02": "coal/2023/02/Monthly_Gas_Report_Feb_23_upload.pdf",
    "2023-03": "coal/2023/04/Monthly_Gas_Report_March_23_upload.pdf",
    "2023-04": "coal/2023/04/Monthly_Gas_Report_April_23.pdf",
    "2023-05": "coal/2023/05/Monthly_Gas_Report_May_23.pdf",
    "2023-06": "coal/2023/06/Monthly_Gas_Report_June_23.pdf",
    "2023-07": "coal/2023/07/Monthly_Gas_Report_July_23.pdf",
    "2023-08": "coal/2023/08/Monthly_Gas_Report_August_23.pdf",
    "2023-10": "coal/2023/10/Monthly_Gas_Report_October_23.pdf",
}


def cea_file(y, m, url):
    return RAW / "cea" / f"{y}-{m:02d}_{url.rsplit('/', 1)[-1]}"


# ----------------------------------------------------------------------------------------------- CEA parsing

STATES = ["ANDHRA PRADESH", "ARUNACHAL PRADESH", "ASSAM", "BIHAR", "CHHATTISGARH", "DELHI", "GOA", "GUJARAT", "HARYANA",
          "HIMACHAL PRADESH", "JAMMU & KASHMIR", "JAMMU AND KASHMIR", "JHARKHAND", "KARNATAKA", "KERALA", "MADHYA PRADESH",
          "MAHARASHTRA", "MANIPUR", "MEGHALAYA", "MIZORAM", "NAGALAND", "ODISHA", "ORISSA", "PUDUCHERRY", "PONDICHERRY",
          "PUNJAB", "RAJASTHAN", "SIKKIM", "TAMIL NADU", "TAMILNADU", "TELANGANA", "TRIPURA", "UTTAR PRADESH", "UTTARAKHAND",
          "UTTRAKHAND", "WEST BENGAL", "ANDAMAN & NICOBAR", "ANDAMAN AND NICOBAR"]
STATE_RE = "|".join(sorted((re.escape(s).replace(r"\ ", r"\s*") for s in STATES), key=len, reverse=True))
ROW_RE = re.compile(rf"^(\d{{1,3}})\s*\.?\s+(.+?)\s+(\d+(?:\.\d+)?)\s*({STATE_RE})\b\s*(?:([PIpi])\s+)?(.*)$", re.I)
GRAND_RE = re.compile(r"^GRAND\s*TOTAL\S*\s+(.*)$", re.I)
SECTION_RE = re.compile(r"^\(?[A-C]\)?\s*(CENTRAL|STATE|PVT|PRIVATE)|^PVT\b|^CENTRAL SECTOR|^STATE SECTOR|^Plants\s+connected", re.I)
NUM_TOK = re.compile(r"^-?\d+(?:\.\d+)?$|^-$|^DNR$|^NA$|^N\.A\.?$", re.I)
# columns after state/P-I, by count of numeric cells
LAYOUTS = {
    7: ["gen_mu", "gas_required_mmscmd", "gas_allotted_mmscmd", "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl", "gen_loss_mu"],
    17: ["gen_mu", "allotted_apm", "allotted_nonapm", "allotted_kgd6", "a_rlng_lt", "a_rlng_ebid", "gas_allotted_mmscmd",
         "consumed_apm", "consumed_nonapm", "consumed_kgd6", "c_rlng_lt", "c_rlng_ebid", "consumed_rlng_spot",
         "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl", "gen_loss_mu"],
    13: ["gen_mu", "allotted_apm", "allotted_kgd6", "allotted_domestic", "allotted_rlng", "consumed_apm", "consumed_kgd6",
         "consumed_domestic", "c_rlng_lt", "consumed_rlng_spot", "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl"],
    15: ["gen_mu", "allotted_apm", "allotted_nonapm", "a_rlng_lt", "allotted_kgd6", "gas_allotted_mmscmd",
         "consumed_apm", "consumed_nonapm", "consumed_kgd6", "c_rlng_lt", "consumed_rlng_spot", "gas_consumed_mmscmd",
         "naphtha_kl", "hsd_kl", "gen_loss_mu"],
    16: ["gen_mu", "allotted_apm", "allotted_nonapm", "allotted_kgd6", "a_rlng_lt", "a_rlng_ebid", "gas_allotted_mmscmd",
         "consumed_apm", "consumed_nonapm", "consumed_kgd6", "c_rlng_lt", "c_rlng_ebid", "consumed_rlng_spot",
         "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl"],
    9: ["gen_mu", "consumed_apm", "consumed_kgd6", "consumed_domestic", "c_rlng_lt", "consumed_rlng_spot",
        "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl"],
    10: ["gen_mu", "consumed_apm", "consumed_kgd6", "consumed_domestic", "c_rlng_lt", "consumed_rlng_spot", "consumed_rlng",
         "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl"],
    14: ["gen_mu", "allotted_apm", "allotted_kgd6", "allotted_domestic", "allotted_rlng", "consumed_apm", "consumed_kgd6",
         "consumed_domestic", "c_rlng_lt", "consumed_rlng_spot", "consumed_rlng", "gas_consumed_mmscmd", "naphtha_kl", "hsd_kl"],
}


def page_lines(page):
    """Text lines of a page, dropping the repeated section title that CEA's PDFs print over the first data row
    of continuation pages (e.g. '(A) CENTRAL SECTOR' overlaid on '44 GAUTAMI CCPP ...')."""
    chars = [c for c in page.chars if c.get("upright", True)]
    lines, cur = [], None
    for c in sorted(chars, key=lambda c: c["top"]):
        if cur is None or c["top"] - cur[0]["top"] > 2.5:
            cur = [c]
            lines.append(cur)
        else:
            cur.append(c)
    drop = set()
    for cs in lines:
        groups = {}
        for c in cs:
            groups.setdefault((c["fontname"], round(c["size"], 1)), []).append(c)
        if len(groups) < 2:
            continue
        texts = {k: "".join(c["text"] for c in sorted(v, key=lambda c: c["x0"])).strip() for k, v in groups.items()}
        if any(re.match(r"^\d", t) for t in texts.values()):
            for k, t in texts.items():
                if SECTION_RE.match(t):
                    drop.update(id(c) for c in groups[k])
    if drop:
        page = page.filter(lambda o: id(o) not in drop)
    return (page.extract_text() or "").splitlines()


def split_nums(rest):
    toks = re.sub(r"(\d)([A-Za-z])", r"\1 \2", rest).split()
    vals = []
    for i, t in enumerate(toks):
        t2 = t.replace(",", "")
        if NUM_TOK.match(t2):
            vals.append(float(t2) if re.match(r"^-?\d", t2) else None)
        else:
            return vals, " ".join(toks[i:])
    return vals, ""


IDENTITIES = {
    17: [("gas_allotted_mmscmd", ["allotted_apm", "allotted_nonapm", "allotted_kgd6", "a_rlng_lt", "a_rlng_ebid"]),
         ("gas_consumed_mmscmd", ["consumed_apm", "consumed_nonapm", "consumed_kgd6", "c_rlng_lt", "c_rlng_ebid", "consumed_rlng_spot"])],
    13: [("allotted_domestic", ["allotted_apm", "allotted_kgd6"]), ("consumed_domestic", ["consumed_apm", "consumed_kgd6"]),
         ("gas_consumed_mmscmd", ["consumed_domestic", "c_rlng_lt", "consumed_rlng_spot"])],
    14: [("allotted_domestic", ["allotted_apm", "allotted_kgd6"]), ("consumed_domestic", ["consumed_apm", "consumed_kgd6"]),
         ("consumed_rlng", ["c_rlng_lt", "consumed_rlng_spot"]), ("gas_consumed_mmscmd", ["consumed_domestic", "consumed_rlng"])],
    7: [],
    15: [("gas_allotted_mmscmd", ["allotted_apm", "allotted_nonapm", "a_rlng_lt", "allotted_kgd6"]),
         ("gas_consumed_mmscmd", ["consumed_apm", "consumed_nonapm", "consumed_kgd6", "c_rlng_lt", "consumed_rlng_spot"])],
    10: [("consumed_domestic", ["consumed_apm", "consumed_kgd6"]), ("consumed_rlng", ["c_rlng_lt", "consumed_rlng_spot"]),
         ("gas_consumed_mmscmd", ["consumed_domestic", "consumed_rlng"])],
}
IDENTITIES[16] = IDENTITIES[17]
IDENTITIES[9] = [("consumed_domestic", ["consumed_apm", "consumed_kgd6"]),
                 ("gas_consumed_mmscmd", ["consumed_domestic", "c_rlng_lt", "consumed_rlng_spot"])]


def score(rec, lay):
    sc = 0
    for tot, parts in IDENTITIES[lay]:
        total = rec.get(tot)
        if total is None:
            continue
        sc += 1 if abs(total - sum(rec.get(k) or 0 for k in parts)) <= 0.011 * len(parts) + 0.001 else -1
    return sc


def realign(v, lay):
    """A row with one or two cells blank in the PDF: put the gaps where the row's own totals add up best.
    Ties: gaps in the alternate-fuel / gen-loss columns first, then in the allotment block, then elsewhere."""
    k = lay - len(v)
    cols = LAYOUTS[lay]
    rank = [0 if c in ("naphtha_kl", "hsd_kl", "gen_loss_mu") else 1 if c.startswith(("allotted", "a_", "gas_allotted")) else 2
            for c in cols]
    best, best_key = v + [None] * k, None
    for gaps in combinations(range(lay), k):
        it = iter(v)
        cand = [None if p in gaps else next(it) for p in range(lay)]
        key = (-score(dict(zip(cols, cand)), lay), sum(rank[p] for p in gaps), gaps)
        if best_key is None or key < best_key:
            best, best_key = cand, key
    return best


def fill_totals(rec, lay):
    """A missing total is rebuilt from its parts."""
    for tot, parts in IDENTITIES[lay]:
        if rec.get(tot) is None and any(rec.get(k) is not None for k in parts):
            rec[tot] = round(sum(rec.get(k) or 0 for k in parts), 4)
    return rec


def derive(rec, lay=None):
    """Fill the comparable columns (domestic / RLNG / total) from whatever split a layout has."""
    def s(*ks):
        return None if all(rec.get(k) is None for k in ks) else round(sum(rec.get(k) or 0 for k in ks), 4)
    for key, parts in (("allotted_domestic", ("allotted_apm", "allotted_nonapm", "allotted_kgd6")),
                       ("allotted_rlng", ("a_rlng_lt", "a_rlng_ebid")),
                       ("consumed_domestic", ("consumed_apm", "consumed_nonapm", "consumed_kgd6")),
                       ("consumed_rlng", ("c_rlng_lt", "c_rlng_ebid", "consumed_rlng_spot")),
                       ("gas_allotted_mmscmd", ("allotted_domestic", "allotted_rlng"))):
        if rec.get(key) is None:
            rec[key] = s(*parts)
    for k in ("a_rlng_lt", "a_rlng_ebid", "c_rlng_lt", "c_rlng_ebid"):
        rec.pop(k, None)
    return rec


class ScannedPDF(Exception):
    pass


def subseq(needle, hay):
    it = iter(hay)
    return all(ch in it for ch in needle)


def parse_cea_pdf(path):
    """-> (layout, plant records, grand-total record, report month text)."""
    lines = []
    with pdfplumber.open(path) as pdf:
        if not any(p.chars for p in pdf.pages):
            raise ScannedPDF(f"{path.name}: scanned image without a text layer, not parsed")
        for p in pdf.pages:
            lines += page_lines(p)
    title = next((l.strip() for l in lines[:12] if re.search(r"(19|20)\d\d|\b[A-Za-z]{3,9}[-,' ]+\d{2}\b", l) and len(l) < 40), "")
    raw, grand, sector, supply = [], None, None, None
    i = 0
    while i < len(lines):
        l = lines[i].strip()
        i += 1
        low = l.lower()
        if re.match(r"^\(?[a-c]?\)?\s*central\s+sector", low):
            sector = "Central"
        elif re.match(r"^\(?[a-c]?\)?\s*state\s+sector", low):
            sector = "State"
        elif re.match(r"^\(?[a-c]?\)?\s*(pvt|private)", low):
            sector = "Private"
        if low.startswith("pl") and subseq("isolatedfield", low.replace(" ", "")):
            supply = "I"                       # also catches the doubled-glyph 'PPllaannttss ... gIsaosl agtreidd fields'
        elif low.startswith("pl") and "gas grid" in low:
            supply = "P"
        gm = GRAND_RE.match(l)
        if gm:
            vals, _ = split_nums(gm[1])
            if len(vals) >= 7:
                grand = vals
            continue
        m = ROW_RE.match(l)
        if not m:
            continue
        vals, remark = split_nums(m[6])
        # numbers alone on the next line may be the tail of a row broken across two lines
        extra = []
        if i < len(lines) and not ROW_RE.match(lines[i].strip()):
            more, rem2 = split_nums(lines[i].strip())
            if more and not rem2:
                extra = more
        name = re.sub(r"\s+", " ", m[2]).strip(" .*#@")
        # a name wrapped onto the next line ('PHASE-' / 'II', '(LANCO' / 'TANJORE)')
        if i < len(lines) and re.fullmatch(r"[A-Z][A-Z0-9 ().\-]{0,25}", lines[i].strip()) and not SECTION_RE.match(lines[i].strip()) \
                and (name.count("(") > name.count(")") or name.endswith("-")):
            name = f"{name} {lines[i].strip()}".replace("- ", "-")
            i += 1
        raw.append({"plant": name, "sector": sector, "supply": (m[5] or supply or "").upper() or None,
                    "state": re.sub(r"\s+", " ", m[4]).upper(), "capacity_mw": float(m[3]), "vals": vals, "extra": extra,
                    "remarks": remark or None})
    counts = {}
    for r in raw:
        counts[len(r["vals"])] = counts.get(len(r["vals"]), 0) + 1
    if grand and len(grand) - 1 in LAYOUTS:
        lay = len(grand) - 1
    else:
        lay = max((c for c in counts if c in LAYOUTS), key=lambda c: counts[c], default=None)
    if lay is None:
        raise ValueError(f"{path.name}: unknown layout, value counts {counts}")
    recs, bad = [], []
    for r in raw:
        v, extra = r.pop("vals"), r.pop("extra")
        if len(v) < lay and len(v) + len(extra) == lay:
            v = v + extra
        if lay - 2 <= len(v) < lay:
            bad.append(f"{r['plant']} ({lay - len(v)} blank, realigned)")
            v = realign(v, lay)
        elif len(v) != lay:
            bad.append(f"{r['plant']} ({len(v)} of {lay} cells)")
            v = (v + [None] * lay)[:lay]
        r.update(zip(LAYOUTS[lay], v))
        recs.append(derive(fill_totals(r, lay), lay))
    g = derive(dict(zip(["capacity_mw"] + LAYOUTS[lay], grand)), lay) if grand else None
    return lay, recs, g, title, bad


def store_plants(month, recs, src):
    con = connect()
    con.executescript(SCHEMA)
    seen = {}
    rows = []
    for r in recs:
        name = r["plant"]
        if name in seen:                      # same name twice in one report (e.g. grid and isolated blocks)
            seen[name] += 1
            name = f"{name} #{seen[name]}"
        else:
            seen[name] = 1
        r = dict(r, plant=name, month=month, source_file=src)
        rows.append(tuple(r.get(c) for c in COLS))
    with con:
        con.execute("DELETE FROM gas_plants_monthly WHERE month=?", (month,))
        con.executemany(f"INSERT OR REPLACE INTO gas_plants_monthly ({','.join(COLS)}) VALUES ({','.join('?' * len(COLS))})", rows)
    con.close()
    return len(rows)


def national_rows(month, g):
    out = []
    for series, key, unit in (("gas_cea_m:Consumed", "gas_consumed_mmscmd", "MMSCMD"),
                              ("gas_cea_m:Allotted", "gas_allotted_mmscmd", "MMSCMD"),
                              ("gas_cea_m:Required", "gas_required_mmscmd", "MMSCMD"),
                              ("gas_cea_m:Consumed_Domestic", "consumed_domestic", "MMSCMD"),
                              ("gas_cea_m:Consumed_RLNG", "consumed_rlng", "MMSCMD"),
                              ("gas_cea_m:Allotted_Domestic", "allotted_domestic", "MMSCMD"),
                              ("gas_cea_m:Allotted_RLNG", "allotted_rlng", "MMSCMD"),
                              ("gas_gen_m:Total", "gen_mu", "MU"),
                              ("gas_genloss_m:Total", "gen_loss_mu", "MU"),
                              ("gas_cea_cap_m:Total", "capacity_mw", "MW")):
        if g.get(key) is not None:
            out.append((series, month, g[key], unit, SRC_CEA))
    return out


def cea_month(y, m, refresh=False):
    """Download and load one month's CEA gas report. Returns a status string, or None if none is listed."""
    links = cea_links(y, m, refresh=refresh)
    pdfs = [u for u in links if u.lower().endswith(".pdf")]
    if not pdfs and f"{y}-{m:02d}" in UNLISTED:
        pdfs = ["https://cea.nic.in/wp-content/uploads/" + UNLISTED[f"{y}-{m:02d}"]]
    if not pdfs:
        return None
    # a month listed twice (re-upload): take the file from the latest upload folder .../uploads/coal/YYYY/MM/
    url = max(pdfs, key=lambda u: (re.findall(r"/(20\d\d)/(\d\d)/", u) or [("", "")])[-1])
    path = cea_file(y, m, url)
    if fetch(url, path) is None:
        raise RuntimeError(f"404 {url}")
    for u in links:                            # keep the xlsx twin too when there is one
        if u.lower().endswith((".xls", ".xlsx")):
            fetch(u, cea_file(y, m, u))
    lay, recs, g, title, bad = parse_cea_pdf(path)
    month = f"{y}-{m:02d}"
    n = store_plants(month, recs, path.name)
    if g:
        put(national_rows(month, g))
    tot = sum(r.get("gas_consumed_mmscmd") or 0 for r in recs)
    return (f"{month} layout {lay} plants {n} consumed {g and g.get('gas_consumed_mmscmd')} (sum {tot:.2f}) "
            f"allotted {g and g.get('gas_allotted_mmscmd')} gen {g and g.get('gen_mu')} title '{title}'"
            + (f" irregular rows {bad}" if bad else "") + ("" if g else " NO GRAND TOTAL"))


# ----------------------------------------------------------------------------------------------- CEA annual (xlsx)

def cea_annual_links(y, refresh=False):
    body = fetch(CEA_AJAX, RAW / "cea_ajax_annual" / f"{y}.html", referer=CEA_ANNUAL_REFERER, refresh=refresh,
                 post={"action": "yearly_archive_report", "selYear": str(y), "reportType": "fuel_consumption"})
    return [u for u in dict.fromkeys(re.findall(r'href="(https://cea\.nic\.in/wp-content/uploads/[^"]+\.xlsx?)"',
                                                   (body or b"").decode("utf8", "ignore"), re.I))
            if re.search(r"gas", u.rsplit("/", 1)[-1], re.I)]


def cea_annual(y, refresh=False):
    """Annual gas report xlsx (FY y..y+1) -> gas_plants_annual (yearly-average MMSCMD, annual MU); the columns
    follow the monthly layout of the time (13 cells in 2022-23, 14 in 2024-25)."""
    out = []
    for url in cea_annual_links(y, refresh=refresh):
        path = RAW / "cea_annual" / url.rsplit("/", 1)[-1]
        if fetch(url, path) is None:
            continue
        fy = f"{y}-{(y + 1) % 100:02d}"
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        rows = [list(r) + [None] * 20 for r in ws.iter_rows(values_only=True)]
        h = next(i for i, r in enumerate(rows) if str(r[0] or "").strip().lower().startswith("s. no"))
        pi = next((j for j, v in enumerate(rows[h]) if str(v or "").replace(" ", "").upper().startswith("P/I")), None)
        start = pi + 1 if pi is not None else 4
        tot = max(j for j, v in enumerate(rows[h + 1]) if str(v or "").strip().upper() == "TOTAL")   # the grand TOTAL column
        lay = 14 if tot - start == 11 else 13          # 2024-25 has a separate RLNG total column
        sector = supply = None
        recs, seen = [], {}
        for r in rows:
            lab = " ".join(str(v) for v in r[:2] if v is not None).lower()
            if "connected with gas grid" in lab:
                supply = "P"
            elif "isolated" in lab and "total" not in lab:
                supply = "I"
            for key, sec in (("central sector", "Central"), ("state sector", "State"), ("private sector", "Private")):
                if key in lab:
                    sector = sec
            if not isinstance(r[0], (int, float)) or not r[1] or num(r[2]) is None:
                continue
            name = re.sub(r"\s+", " ", str(r[1])).strip(" .*#")
            seen[name] = seen.get(name, 0) + 1
            if seen[name] > 1:
                name = f"{name} #{seen[name]}"
            rec = {"fy": fy, "plant": name, "sector": sector, "state": str(r[3] or "").strip().upper() or None,
                   "supply": (str(r[pi] or "").strip().upper() or None) if pi is not None else supply,
                   "capacity_mw": num(r[2]), "source_file": path.name}
            rec.update(zip(LAYOUTS[lay], (num(v) for v in r[start:start + lay])))
            rec["consumed_rlng_lt"] = rec.get("c_rlng_lt")
            derive(rec)
            recs.append(tuple(rec.get(c) for c in ACOLS))
        con = connect()
        con.executescript(SCHEMA)
        with con:
            con.execute("DELETE FROM gas_plants_annual WHERE fy=?", (fy,))
            con.executemany(f"INSERT INTO gas_plants_annual ({','.join(ACOLS)}) VALUES ({','.join('?' * len(ACOLS))})", recs)
        con.close()
        out.append(f"annual {fy}: {len(recs)} plants")
    return out


# ----------------------------------------------------------------------------------------------- entry points

def cea(first=FIRST_ARCHIVE, refresh_recent=4):
    """Walk the CEA archive from `first` to last month; the last `refresh_recent` months' listings are re-read."""
    t = date.today()
    last = (t.year, t.month - 1) if t.month > 1 else (t.year - 1, 12)
    allm = list(months(first, last))
    recent = set(allm[-refresh_recent:])
    msgs, found, failed = [], [], []
    for y, m in allm:
        try:
            msg = cea_month(y, m, refresh=(y, m) in recent)
        except ScannedPDF as e:                  # 2016-09 (OCR of the scan is unusable)
            msgs.append(f"{y}-{m:02d} skipped: {e}")
            continue
        except Exception as e:  # noqa: BLE001 - keep walking, report at the end
            failed.append(f"{y}-{m:02d}: {e}")
            continue
        if msg:
            found.append(f"{y}-{m:02d}")
            msgs.append(msg)
    return msgs, found, failed


def run():
    """Refresh (collect.py): PPAC files, the last few CEA archive months, the latest annual report."""
    out = []
    try:
        out.append(ppac(refresh=True))
    except Exception as e:  # noqa: BLE001
        log("gas_power", "error", f"PPAC: {e}")
        out.append(f"PPAC failed: {e}")
    t = date.today()
    first = (t.year, t.month - 4) if t.month > 4 else (t.year - 1, t.month + 8)
    msgs, found, failed = cea(first=first, refresh_recent=4)
    out += msgs + [f"CEA failed {f}" for f in failed]
    try:
        out += cea_annual(t.year - 1 if t.month < 4 else t.year, refresh=True)
    except Exception as e:  # noqa: BLE001
        out.append(f"annual failed: {e}")
    # cea.nic.in was rebuilt (Tomcat app, /fuel-reports now redirects to /loginPage, admin-ajax 403 even via
    # the proxy), so the listing is gone. Known-dead source: report as info, not as a scraper failure.
    if failed and all("403" in f for f in failed):
        log("gas_power", "info", f"CEA archive unavailable (site rebuilt); PPAC ok; skipped {len(failed)} months")
    else:
        log("gas_power", "error" if failed else "ok", f"CEA months {found or 'none new'}; failed {failed}")
    return out


def backfill():
    print(ppac(refresh=True))
    msgs, found, failed = cea(first=FIRST_ARCHIVE, refresh_recent=4)
    for m in msgs:
        print(m)
    for f in failed:
        print("FAILED", f)
    print(f"CEA: {len(found)} months with a gas report: {found[0] if found else ''}..{found[-1] if found else ''}")
    for y in range(2010, date.today().year + 1):
        for m in cea_annual(y):
            print(m)
    log("gas_power", "error" if failed else "ok", f"backfill: CEA {len(found)} months, failed {failed}")


if __name__ == "__main__":
    if "--backfill" in sys.argv:
        backfill()
    else:
        for line in run():
            print(line)
